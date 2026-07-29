"""Renders a slide deck out as a PowerPoint file."""

from __future__ import annotations

import logging
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from backend.models.artifacts import Slide, SlidesModel

logger = logging.getLogger(__name__)

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
BLANK_LAYOUT = 6

PRIMARY = RGBColor(0x1E, 0x1B, 0x4B)
ACCENT = RGBColor(0x63, 0x66, 0xF1)
BODY = RGBColor(0x36, 0x36, 0x36)
SUBTITLE = RGBColor(0xCB, 0xD5, 0xE1)
HEADER_FILL = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


class DeckError(ValueError):
    """Raised when a deck is missing something we need to render it."""


class SlidesPptxRenderer:
    """Builds a 16:9 deck: a title slide, then one slide for each entry."""

    def render(self, deck: SlidesModel, destination: Path) -> Path:
        """Write `deck` out to `destination`, and hand that path back."""
        self._validate(deck)
        logger.info("Rendering %d slides", len(deck.slides))

        presentation = Presentation()
        presentation.slide_width = SLIDE_WIDTH
        presentation.slide_height = SLIDE_HEIGHT

        self._title_slide(presentation, deck)
        for index, entry in enumerate(deck.slides, start=1):
            self._content_slide(presentation, entry, index)

        destination.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(str(destination))
        return destination

    @staticmethod
    def _validate(deck: SlidesModel) -> None:
        if not deck.title:
            raise DeckError("The deck has no title")
        if not deck.slides:
            raise DeckError("The deck has no slides")
        for index, slide in enumerate(deck.slides):
            if not slide.heading:
                raise DeckError(f"Slide {index + 1} has no heading")

    def _title_slide(self, presentation: Presentation, deck: SlidesModel) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK_LAYOUT])

        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = PRIMARY

        self._text(slide, deck.title, top=2.25, size=44, colour=WHITE, bold=True)
        self._text(
            slide,
            f"Audience: {deck.audience_level or 'General'}",
            top=3.75, height=0.6, size=20, colour=SUBTITLE,
        )
        self._bar(slide, top=7.1, height=0.4, colour=ACCENT)

    def _content_slide(self, presentation: Presentation, entry: Slide, number: int) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK_LAYOUT])

        self._bar(slide, top=0, height=1.0, colour=HEADER_FILL)
        self._text(slide, entry.heading, top=0.1, left=0.5, width=12.333, height=0.8,
                   size=32, colour=PRIMARY, bold=True)
        self._bar(slide, top=1.1, height=0.03, left=1, width=11.333, colour=ACCENT)

        if entry.main_idea:
            self._text(slide, entry.main_idea, top=1.4, left=1.5, width=10.333, height=0.8,
                       size=20, colour=RGBColor(0x43, 0x38, 0xCA), italic=True)

        if entry.bullet_points:
            self._bullets(slide, entry.bullet_points)

        if entry.speaker_notes:
            try:
                slide.notes_slide.notes_text_frame.text = entry.speaker_notes
            except Exception as error:
                logger.debug("Could not attach notes to slide %d: %s", number, error)

    @staticmethod
    def _text(
        slide,
        text: str,
        *,
        top: float,
        left: float = 0,
        width: float = 13.333,
        height: float = 1.5,
        size: int,
        colour: RGBColor,
        bold: bool = False,
        italic: bool = False,
    ) -> None:
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        box.text_frame.word_wrap = True

        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = str(text)
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.italic = italic
        paragraph.font.color.rgb = colour
        paragraph.alignment = PP_ALIGN.CENTER

    @staticmethod
    def _bullets(slide, points: list[str]) -> None:
        box = slide.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(11.333), Inches(4.5))
        frame = box.text_frame
        frame.word_wrap = True

        for index, point in enumerate(points):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = f"{index + 1}. {point}"
            paragraph.font.size = Pt(20)
            paragraph.font.color.rgb = BODY
            paragraph.space_after = Pt(16)

    @staticmethod
    def _bar(
        slide,
        *,
        top: float,
        height: float,
        colour: RGBColor,
        left: float = 0,
        width: float = 13.333,
    ) -> None:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = colour
        shape.line.fill.background()
